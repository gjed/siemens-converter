"""Tests for static_reader.py — parsing static condominium data XLSX."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from siemens_converter.static_reader import parse_static_data

FIXTURE = Path(__file__).parent / "fixtures" / "static_data_TEST.xlsx"


def test_parse_valid_file():
    sd = parse_static_data(FIXTURE)
    assert len(sd.apartments) == 10
    assert len(sd.millesimali) == 10
    assert len(sd.costs) == 6
    assert len(sd.meters) == 3
    assert len(sd.previous_readings) == 10


def test_apartments_populated():
    sd = parse_static_data(FIXTURE)
    apt1 = sd.apartments[0]
    assert apt1.apartment_number == 1
    assert apt1.proprietario == "Rossi Mario"
    assert apt1.inquilino == "Bianchi Anna"


def test_apartment_empty_inquilino():
    sd = parse_static_data(FIXTURE)
    apt10 = sd.apartments[9]
    assert apt10.apartment_number == 10
    assert apt10.inquilino == ""


def test_millesimali_populated():
    sd = parse_static_data(FIXTURE)
    m1 = sd.millesimali[0]
    assert m1.apartment_number == 1
    assert m1.subalterno == 20
    assert m1.heat_energy_kwh == 1344.03
    assert m1.water_energy_kwh == 1101.36


def test_costs_populated():
    sd = parse_static_data(FIXTURE)
    cost_dict = {c.label: c.amount for c in sd.costs}
    assert cost_dict["Energia elettrica"] == 263.44
    assert cost_dict["Gas metano"] == 2141.0
    assert cost_dict["Acqua condominio"] == 0.0


def test_meters_populated():
    sd = parse_static_data(FIXTURE)
    # Fixture now stores uppercase names matching the template display style
    m = {mr.name.upper(): mr for mr in sd.meters}
    assert m["ENERGIA ELETTRICA CT"].initial == 0
    assert m["ENERGIA ELETTRICA CT"].final == 1086
    assert m["GAS METANO CT"].final == 1718


def test_previous_readings_populated():
    sd = parse_static_data(FIXTURE)
    pr1 = sd.previous_readings[0]
    assert pr1.apartment_number == 1
    assert pr1.heat_kwh == 4048
    assert pr1.water_m3 == 25.5
    assert pr1.cold_water_m3 == 10.2


def test_missing_sheet_raises_error(tmp_path: Path):
    """File with a required sheet removed should raise ValueError."""
    wb = openpyxl.load_workbook(FIXTURE)
    del wb["Contatori"]
    bad_path = tmp_path / "missing_sheet.xlsx"
    wb.save(bad_path)

    with pytest.raises(ValueError, match="Contatori"):
        parse_static_data(bad_path)


def test_missing_cost_label_defaults_to_zero(tmp_path: Path):
    """Costi sheet missing 'Acqua condominio' row should default to 0.0."""
    wb = openpyxl.load_workbook(FIXTURE)
    ws = wb["Costi"]
    # New layout: row1=title, row2=header, row3=Energia el., row4=Gas metano,
    #             row5=Acqua condominio, ...
    ws.delete_rows(5)  # remove "Acqua condominio"
    patched = tmp_path / "missing_cost.xlsx"
    wb.save(patched)

    sd = parse_static_data(patched)
    cost_dict = {c.label: c.amount for c in sd.costs}
    assert cost_dict["Acqua condominio"] == 0.0
    # Other costs should still be present
    assert cost_dict["Gas metano"] == 2141.0


def test_fewer_apartment_rows(tmp_path: Path):
    """Inquilini with only 3 data rows should produce 3 apartments (not error)."""
    wb = openpyxl.load_workbook(FIXTURE)
    ws = wb["Inquilini"]
    # New layout: row1=title, row2=header, rows3-12=apts 1-10.
    # Keep only rows 3,4,5 (apts 1,2,3) — delete rows 6-12 (7 rows from row 6).
    ws.delete_rows(6, 7)
    patched = tmp_path / "fewer_apts.xlsx"
    wb.save(patched)

    sd = parse_static_data(patched)
    assert len(sd.apartments) == 3
    assert sd.apartments[0].apartment_number == 1
    assert sd.apartments[2].apartment_number == 3
