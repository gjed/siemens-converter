"""Tests for CLI entry point."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

FIXTURE_DIR = Path(__file__).parent / "fixtures"
FC_REPORT = FIXTURE_DIR / "FC_report_TEST_9999_2026-01-15.xls"
STATIC_DATA = FIXTURE_DIR / "static_data_TEST.xlsx"


def test_main_no_args(monkeypatch):
    monkeypatch.setattr(sys, "argv", ["siemens-converter"])
    from siemens_converter.__main__ import main

    with pytest.raises(SystemExit):
        main()


def test_main_missing_file(monkeypatch):
    monkeypatch.setattr(sys, "argv", ["siemens-converter", "nonexistent.xls"])
    from siemens_converter.__main__ import main

    with pytest.raises(SystemExit):
        main()


def test_main_success_single_file(monkeypatch, tmp_path):
    """Single FC_report argument — backward compatible."""
    input_file = tmp_path / FC_REPORT.name
    shutil.copy(FC_REPORT, input_file)

    monkeypatch.setattr(sys, "argv", ["siemens-converter", str(input_file)])
    from siemens_converter.__main__ import main

    main()

    outputs = list(tmp_path.glob("Riparto_*.xlsx"))
    assert len(outputs) == 1

    # Inquilini tenant column should be empty (no static data)
    wb = openpyxl.load_workbook(outputs[0])
    ws = wb["Inquilini"]
    assert ws["B2"].value is None


def test_main_success_two_files(monkeypatch, tmp_path):
    """Both FC_report and static data arguments."""
    input_file = tmp_path / FC_REPORT.name
    shutil.copy(FC_REPORT, input_file)
    static_file = tmp_path / STATIC_DATA.name
    shutil.copy(STATIC_DATA, static_file)

    monkeypatch.setattr(
        sys,
        "argv",
        ["siemens-converter", str(input_file), str(static_file)],
    )
    from siemens_converter.__main__ import main

    main()

    outputs = list(tmp_path.glob("Riparto_*.xlsx"))
    assert len(outputs) == 1

    # Inquilini tenant column should be populated
    wb = openpyxl.load_workbook(outputs[0])
    ws = wb["Inquilini"]
    assert ws["B2"].value == "Bianchi Anna"  # Apt 1 tenant from static data

    # Millesimali should be populated
    ws_mill = wb["Tabelle millesimali"]
    assert ws_mill.cell(row=4, column=2).value == 20  # subalterno apt 1

    # Costs should be populated
    ws_tab = wb["Tabella_2026"]
    assert ws_tab.cell(row=3, column=5).value == 263.44  # Energia elettrica


def test_main_static_file_not_found(monkeypatch, tmp_path):
    """Static data file that doesn't exist should cause exit."""
    input_file = tmp_path / FC_REPORT.name
    shutil.copy(FC_REPORT, input_file)

    monkeypatch.setattr(
        sys,
        "argv",
        ["siemens-converter", str(input_file), str(tmp_path / "nonexistent.xlsx")],
    )
    from siemens_converter.__main__ import main

    with pytest.raises(SystemExit):
        main()


def test_usage_message_mentions_static(monkeypatch, capsys):
    """Usage message should mention the optional static_data argument."""
    monkeypatch.setattr(sys, "argv", ["siemens-converter"])
    from siemens_converter.__main__ import main

    with pytest.raises(SystemExit):
        main()

    captured = capsys.readouterr()
    assert "static_data.xlsx" in captured.out
