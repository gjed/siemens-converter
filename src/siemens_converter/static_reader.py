"""Parse static condominium data from a user-populated XLSX file."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from siemens_converter.models import (
    ApartmentInfo,
    CostItem,
    MeterReading,
    Millesimali,
    PreviousReading,
    StaticData,
)

REQUIRED_SHEETS = frozenset(
    {"Inquilini", "Millesimali", "Costi", "Contatori", "Letture_precedenti"}
)

# Labels expected in the Costi sheet (order matches Tabella_2026 E3-E8)
COST_LABELS = (
    "Energia elettrica",
    "Gas metano",
    "Acqua condominio",
    "Conduzione e manutenzione",
    "Contabilizzazione",
    "Acqua sanitaria manutenzione",
)


def parse_static_data(path: str | Path) -> StaticData:
    """Read a user-populated static data XLSX and return a StaticData object.

    Raises ValueError if required sheets are missing.
    Missing cost rows default to 0.0; fewer apartment rows are silently accepted.
    """
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)

    missing = REQUIRED_SHEETS - set(wb.sheetnames)
    if missing:
        raise ValueError(
            f"Static data file is missing required sheet(s): {', '.join(sorted(missing))}"
        )

    apartments = _read_inquilini(wb["Inquilini"])
    millesimali = _read_millesimali(wb["Millesimali"])
    costs = _read_costi(wb["Costi"])
    meters = _read_contatori(wb["Contatori"])
    previous_readings = _read_letture_precedenti(wb["Letture_precedenti"])

    return StaticData(
        apartments=apartments,
        millesimali=millesimali,
        costs=costs,
        meters=meters,
        previous_readings=previous_readings,
    )


_DATA_START_ROW = 3  # row 1 = section title, row 2 = column headers, row 3 = first data


def _read_inquilini(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[ApartmentInfo]:
    """Read Inquilini sheet: App.number, Proprietario, Inquilino."""
    result: list[ApartmentInfo] = []
    for row in ws.iter_rows(min_row=_DATA_START_ROW, values_only=True):
        apt_num = row[0] if len(row) > 0 else None
        if apt_num is None:
            break
        try:
            apt_num = int(apt_num)
        except (ValueError, TypeError):
            break
        result.append(
            ApartmentInfo(
                apartment_number=apt_num,
                proprietario=str(row[1]) if len(row) > 1 and row[1] else "",
                inquilino=str(row[2]) if len(row) > 2 and row[2] else "",
            )
        )
    return result


def _read_millesimali(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[Millesimali]:
    """Read Millesimali sheet: App., Subalterno, Energia riscaldamento kWh, Energia ACS kWh.

    Column layout: App(1), Subalterno(2), Heat kWh(3), Mill.risc formula(4), ACS kWh(5), Mill.ACS formula(6).
    Non-integer rows (e.g. TOTALE) are skipped rather than stopping iteration.
    """
    result: list[Millesimali] = []
    for row in ws.iter_rows(min_row=_DATA_START_ROW, values_only=True):
        apt_num = row[0] if len(row) > 0 else None
        if apt_num is None:
            continue  # skip empty rows
        try:
            apt_num = int(apt_num)
        except (ValueError, TypeError):
            continue  # skip TOTALE and other non-data rows
        result.append(
            Millesimali(
                apartment_number=apt_num,
                subalterno=int(row[1]) if len(row) > 1 and row[1] is not None else 0,
                heat_energy_kwh=float(row[2])
                if len(row) > 2 and row[2] is not None
                else 0.0,
                # Column 5 (index 4) = ACS kWh; column 4 (index 3) may be a formula cell
                water_energy_kwh=float(row[4])
                if len(row) > 4 and row[4] is not None
                else 0.0,
            )
        )
    return result


def _read_costi(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[CostItem]:
    """Read Costi sheet: Voce, Importo.  Missing expected labels default to 0.0."""
    found: dict[str, float] = {}
    for row in ws.iter_rows(min_row=_DATA_START_ROW, values_only=True):
        label = row[0] if len(row) > 0 else None
        if label is None:
            break
        # Skip the TOTALE row or any non-cost label
        if str(label).upper() == "TOTALE":
            continue
        amount_raw = row[1] if len(row) > 1 else None
        amount = float(amount_raw) if amount_raw is not None else 0.0
        found[str(label)] = amount

    result: list[CostItem] = []
    for label in COST_LABELS:
        result.append(CostItem(label=label, amount=found.get(label, 0.0)))
    return result


def _read_contatori(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[MeterReading]:
    """Read Contatori sheet: Contatore, Unità, Lettura iniziale, Lettura finale."""
    result: list[MeterReading] = []
    for row in ws.iter_rows(min_row=_DATA_START_ROW, values_only=True):
        name = row[0] if len(row) > 0 else None
        if name is None:
            break
        result.append(
            MeterReading(
                name=str(name),
                unit=str(row[1]) if len(row) > 1 and row[1] else "",
                initial=float(row[2]) if len(row) > 2 and row[2] is not None else 0.0,
                final=float(row[3]) if len(row) > 3 and row[3] is not None else 0.0,
            )
        )
    return result


def _read_letture_precedenti(
    ws: openpyxl.worksheet.worksheet.Worksheet,
) -> list[PreviousReading]:
    """Read Letture_precedenti sheet: App., Riscaldamento kWh, ACS m3, AFS m3."""
    result: list[PreviousReading] = []
    for row in ws.iter_rows(min_row=_DATA_START_ROW, values_only=True):
        apt_num = row[0] if len(row) > 0 else None
        if apt_num is None:
            break
        try:
            apt_num = int(apt_num)
        except (ValueError, TypeError):
            break
        result.append(
            PreviousReading(
                apartment_number=apt_num,
                heat_kwh=int(row[1]) if len(row) > 1 and row[1] is not None else 0,
                water_m3=float(row[2]) if len(row) > 2 and row[2] is not None else 0.0,
                cold_water_m3=float(row[3])
                if len(row) > 3 and row[3] is not None
                else 0.0,
            )
        )
    return result
